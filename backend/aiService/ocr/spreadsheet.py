"""
엑셀 입력 (기획안 5절 "이메일·엑셀·PDF 등 비정형 선적 서류").

## 왜 셀을 좌표로 바꾸는가

`FieldParser` 는 좌표와 앵커로 값을 찾는다. 엑셀에는 픽셀 좌표가 없고 행·열만
있으므로, 그리드를 그대로 평면에 펼쳐 BBox 로 만든다. 그러면 이미지·PDF·엑셀이
모두 같은 `OCRResult` 로 모이고 파서가 하나로 유지된다.

엑셀 전용 매핑 경로를 따로 두는 대안도 있었지만, 그러면 "B/L 번호를 어떻게
찾는가"가 입력 형식마다 갈린다. 앵커 테이블을 고칠 때 네 곳을 고쳐야 하고,
그중 하나를 빠뜨리면 특정 입력에서만 필드가 비는 버그가 된다.

## 좌표가 실제 지면을 뜻하지 않는다

여기서 만든 좌표는 **그리드 상의 상대 위치일 뿐** 종이 위 위치가 아니다.
`FieldParser` 의 구역(REGIONS) 판정은 B/L 서식의 실제 배치를 전제하므로
엑셀에서는 대체로 맞지 않는다. 값을 찾는 것은 앵커 쪽이다 — 엑셀 B/L 은
"A열에 라벨, B열에 값" 형태가 대부분이고 그건 앵커가 잘 잡는다.

## 신뢰도

1.0 이다. 셀에 적힌 글자를 그대로 읽었을 뿐 추측하지 않았다. OCR 의 1.0 과는
의미가 다르며, 그 구분은 `OCRResult.source` 가 진다.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, List, Optional, Tuple

from .types import BBox, OCRResult

# 그리드를 펼칠 캔버스. 라벨 데이터셋 해상도와 맞춘다 — 파서의 구역 비율이
# 그 기준으로 교정돼 있어, 다른 값을 쓰면 입력 형식마다 비율이 달라진다.
_CANVAS_WIDTH = 1654
_CANVAS_HEIGHT = 2340

# 캔버스를 나눌 최소 격자 수. 실제 표가 이보다 작아도 이 값으로 나눈다.
# 3×2 짜리 표가 캔버스를 가득 채우면 셀 하나가 지면 절반이 되어, 앵커의
# "같은 줄" 판정(줄 높이 기준)이 무의미해진다.
_MIN_COLUMNS = 8
_MIN_ROWS = 30

# 셀 사이 여백(칸 너비 대비 비율). 0 으로 두면 옆 칸의 x_min 과 이 칸의
# x_max 가 같아지는데, 앵커 탐색은 값이 라벨보다 **엄격히** 오른쪽에 있을 것을
# 요구하므로(`b.x_min > anchor.x_max`) 딱 붙은 두 칸은 라벨-값 관계로 인식되지
# 않는다. 실제 종이에서도 글자가 칸을 가득 채우지는 않는다.
_CELL_GUTTER_RATIO = 0.06

# 훑을 최대 범위. 사용자가 올린 엑셀에 수천 행짜리 빈 시트가 딸려오는 일이
# 흔하고, 그걸 전부 도는 것은 낭비다.
MAX_ROWS = 400
MAX_COLUMNS = 40


def load_workbook(path: str) -> Any:
    """openpyxl 워크북을 연다. 수식은 계산된 값으로 읽는다."""
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - 의존성 부재 경로
        raise ImportError(
            "엑셀 입력에는 openpyxl 이 필요합니다.\n"
            "설치: pip install openpyxl"
        ) from exc

    # data_only=True 는 **저장 시점에 캐시된 계산 결과**를 준다. 엑셀로 한 번도
    # 열지 않고 만든 파일은 이 값이 없어 None 이 되는데, 그건 수식이 아니라
    # 빈 셀로 보인다. 그래도 수식 문자열("=B2*3")을 값으로 넣는 것보다 낫다 —
    # 그쪽은 B/L 번호 자리에 수식이 박히는 결과가 된다.
    return openpyxl.load_workbook(path, data_only=True, read_only=True)


def from_excel(path: str, sheet: Optional[str] = None) -> OCRResult:
    """엑셀 한 시트를 OCRResult 로 바꾼다.

    sheet 를 주지 않으면 **값이 가장 많은 시트**를 고른다. 첫 시트를 쓰면
    표지·안내 시트가 앞에 있는 파일에서 빈 결과가 나온다.
    """
    workbook = load_workbook(path)
    try:
        worksheet = _pick_sheet(workbook, sheet)
        cells = _read_cells(worksheet)
    finally:
        workbook.close()

    if not cells:
        raise ValueError(f"값이 있는 셀이 없습니다: {path}")

    return OCRResult(
        image_id=Path(path).stem,
        image_width=_CANVAS_WIDTH,
        image_height=_CANVAS_HEIGHT,
        form_type="선하증권",
        bboxes=_to_bboxes(cells),
        source="excel",
    )


def _pick_sheet(workbook: Any, name: Optional[str]) -> Any:
    if name is not None:
        if name not in workbook.sheetnames:
            raise ValueError(
                f"시트를 찾을 수 없습니다: {name} "
                f"(있는 시트: {', '.join(workbook.sheetnames)})"
            )
        return workbook[name]

    best, best_count = None, -1
    for sheet_name in workbook.sheetnames:
        ws = workbook[sheet_name]
        count = sum(1 for _ in _read_cells(ws))
        if count > best_count:
            best, best_count = ws, count
    return best if best is not None else workbook.active


def _read_cells(worksheet: Any) -> List[Tuple[int, int, str]]:
    """(행, 열, 텍스트) 목록. 빈 셀은 뺀다."""
    out: List[Tuple[int, int, str]] = []
    for row in worksheet.iter_rows(max_row=MAX_ROWS, max_col=MAX_COLUMNS):
        for cell in row:
            if cell.value is None:
                continue
            text = str(cell.value).strip()
            if text:
                out.append((cell.row, cell.column, text))
    return out


def _to_bboxes(cells: List[Tuple[int, int, str]]) -> List[BBox]:
    """셀 목록을 캔버스 위 BBox 로 펼친다."""
    max_row = max(max(r for r, _, _ in cells), _MIN_ROWS)
    max_col = max(max(c for _, c, _ in cells), _MIN_COLUMNS)

    cell_w = _CANVAS_WIDTH / max_col
    cell_h = _CANVAS_HEIGHT / max_row

    gutter_x = cell_w * _CELL_GUTTER_RATIO
    gutter_y = cell_h * _CELL_GUTTER_RATIO

    bboxes: List[BBox] = []
    for row, col, text in cells:
        bboxes.append(
            BBox(
                text=text,
                x_min=int((col - 1) * cell_w),
                y_min=int((row - 1) * cell_h),
                x_max=int(col * cell_w - gutter_x),
                y_max=int(row * cell_h - gutter_y),
                # 셀에 적힌 글자를 그대로 읽었다. 추측이 아니다.
                confidence=1.0,
            )
        )
    return bboxes


__all__ = ["from_excel", "load_workbook", "MAX_ROWS", "MAX_COLUMNS"]
